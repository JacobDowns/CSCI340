"""Build the student worksheet for Activity 5 using python-docx."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "activities/week-07-project-brainstorming.docx"
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.7)
section.footer_distance = Inches(0.3)

for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet"):
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(6)
doc.styles["Normal"].font.size = Pt(11)
doc.styles["Normal"].paragraph_format.line_spacing = 1.06
doc.styles["Title"].font.size = Pt(22)
doc.styles["Title"].paragraph_format.space_after = Pt(8)
for style in doc.styles:
    for border in list(style.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
doc.styles["Subtitle"].font.size = Pt(11)
for name, size in (("Heading 1", 15), ("Heading 2", 11)):
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(9)
    doc.styles[name].paragraph_format.space_after = Pt(6)

doc.core_properties.title = "Database project brainstorming"
doc.core_properties.subject = "CSCI 340 Activity 5"
doc.core_properties.author = "CSCI 340"
doc.core_properties.keywords = "database, project, brainstorming, entities, relationships"


def p(text="", style=None):
    return doc.add_paragraph(text, style)


def answer_space(count=2):
    """Expandable blank paragraphs for typed or handwritten answers."""
    for _ in range(count):
        para = p()
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1
        para.paragraph_format.line_spacing_rule = None
        para.paragraph_format.space_after = Pt(6)
        run = para.add_run(" ")
        run.font.size = Pt(14)



def prompt(title, text, count=2):
    p(title, "Heading 2")
    p(text)
    answer_space(count)


footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("CSCI 340  |  Activity 5  |  ")
run.font.size = Pt(9)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

# Page 1: the entire first phase, including an explicit decision.
p("Database project brainstorming", "Title")
p("CSCI 340   Activity 5", "Subtitle")
p("Work with your assigned project group. Use the first 25 minutes to explore two or three possible topics and agree on one. Spend the remainder of class developing your selected idea together.")
group_line = p("Group number:\tDate:")
group_line.paragraph_format.tab_stops.add_tab_stop(Inches(3.6))
p("Contributors")
answer_space(2)
p("First 25 minutes", "Heading 1")
p("For each idea, discuss who would use the database, what problem or activity it would support, and what makes it promising or uncertain. Make room for everyone to contribute. Brief notes are enough.")
for i in range(1, 4):
    p(f"Possible topic {i}" + (" if useful" if i == 3 else ""), "Heading 2")
    answer_space(2)
p("Choose one topic", "Heading 2")
p("By the end of the first 25 minutes, agree on the topic you will develop. Record your choice and why it seems useful, interesting, and manageable.")
answer_space(3)

# Page 2: guided design questions, without a minute-by-minute schedule.
doc.add_page_break()
p("Develop your chosen topic", "Heading 1")
p("Use the remainder of class for these questions and the entity and relationship lists that follow. Build a shared account of the project. Short explanations and lists are sufficient.")
prompt("Purpose and problem", "What is your database for? What specific problem or activity will it address?", 3)
prompt("Intended users", "Who will use it? What information will they enter, update, or retrieve?", 3)
prompt("Questions the database should answer", "List two or three concrete questions users could ask. Include one that brings together information about different kinds of things.", 3)
prompt("Scope", "Describe the smallest useful version. Name one related feature or activity you will leave outside the initial project.", 3)
prompt("Example data", "Where might records come from? You may identify a possible source or describe realistic sample data you could create. Note anything you still need to investigate.", 2)

# Page 3: entity records with substantial handwriting or typing space.
doc.add_page_break()
p("Potential entity types", "Heading 1")
p("Include at least five meaningful entity types; aim for a compact core of five to seven. An entity type describes a kind of thing, such as a member or a loan. An instance is one particular member or one particular loan.")
p("For each candidate, say what one instance represents and suggest a few useful attributes. Each entity should help meet a user need or represent a real rule. Use the space below to list and describe your candidates.")
answer_space(17)
prompt("A distinction worth clarifying", "Choose an entity above. What distinguishes one instance from another? Could names be duplicated, or might a record describe a category instead of an individual object?", 3)

# Page 4: relationships and open questions, no formal diagram requirement.
doc.add_page_break()
p("Potential relationships", "Heading 1")
p("Use a target minimum of four meaningful relationships. Add more if needed to connect your entity types into a coherent model. Describe relationships in words; an informal sketch is optional. A complete ER diagram is not required today.")
p("For example: a member borrows an equipment item. If possible, suggest whether each side can participate once or many times, and whether participation is optional. Mark uncertain rules as questions.")
answer_space(10)
prompt("Modeling questions and assumptions", "What might be difficult or ambiguous to model? Consider repeated events, history, optional relationships, or whether a fact belongs to an entity or an association. If nothing seems ambiguous yet, state an assumption and a case that could challenge it.", 3)
prompt("A decision to investigate", "Choose one question above. What are two possible answers, and how could they change your model? A tentative explanation is enough.", 3)
p("Before leaving class", "Heading 2")
p("Agree on the selected topic, summarize its purpose, and identify your most important open modeling question. Keep one shared group worksheet and make sure every group member can access it. Your topic should be settled today; individual modeling choices can develop as you learn more.")

doc.save(OUTPUT)
print(OUTPUT)
